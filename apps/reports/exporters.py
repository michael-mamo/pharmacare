"""
Generic exporters. Each takes a report instance plus its already-computed
rows/summary and returns an HttpResponse — so a new report gets CSV, Excel
and PDF for free from its column declaration alone.
"""
import csv
import io
import logging
from datetime import datetime
from pathlib import Path

from django.conf import settings
from django.http import HttpResponse
from django.utils import timezone
from django.utils.translation import gettext as _

logger = logging.getLogger(__name__)


def _filename(report, extension):
    stamp = timezone.localtime().strftime("%Y%m%d-%H%M")
    slug = report.slug.replace("_", "-")
    return f"pharmacare-{slug}-{stamp}.{extension}"


def _headers(report):
    return [str(col[0]) for col in report.columns]


#: Characters Excel refuses in a worksheet name. Hitting any of them raises
#: ValueError deep inside openpyxl, which surfaced as a 500 on the
#: "Sales Summary (Daily / Weekly / Monthly / Yearly)" report.
_INVALID_SHEET_CHARS = r'\/*?:[]'


def _safe_sheet_title(title):
    """Excel worksheet names cannot contain \\ / * ? : [ ] and are capped at
    31 characters. Translated titles make this easy to trip, so sanitize
    rather than trusting the report's own title."""
    text = str(title)
    for ch in _INVALID_SHEET_CHARS:
        text = text.replace(ch, "-")
    text = text.strip().strip("'")
    return text[:31] or "Report"


# ---------------------------------------------------------------------------
# CSV
# ---------------------------------------------------------------------------
def export_csv(report, rows, summary, context):
    response = HttpResponse(content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = f'attachment; filename="{_filename(report, "csv")}"'
    # UTF-8 BOM so Excel on Windows opens Amharic text correctly instead of
    # showing mojibake — the default assumption there is a legacy codepage.
    response.write("\ufeff")

    writer = csv.writer(response)
    writer.writerow([context["pharmacy_name"], str(report.title)])
    writer.writerow([_("Generated"), context["generated_at"]])
    if context.get("branch_label"):
        writer.writerow([_("Branch"), context["branch_label"]])
    if context.get("period_label"):
        writer.writerow([_("Period"), context["period_label"]])
    writer.writerow([])

    if summary:
        for label, value in summary:
            writer.writerow([str(label), str(value)])
        writer.writerow([])

    writer.writerow(_headers(report))
    for row in rows:
        writer.writerow([str(cell) for cell in row])
    return response


# ---------------------------------------------------------------------------
# Excel (openpyxl)
# ---------------------------------------------------------------------------
def export_xlsx(report, rows, summary, context):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = _safe_sheet_title(report.title)

    brand_yellow = "FFF1AB15"
    black = "FF000000"

    r = 1
    ws.cell(row=r, column=1, value=context["pharmacy_name"]).font = Font(bold=True, size=14)
    r += 1
    ws.cell(row=r, column=1, value=str(report.title)).font = Font(bold=True, size=12)
    r += 1
    ws.cell(row=r, column=1, value=f"{_('Generated')}: {context['generated_at']}").font = Font(size=9, italic=True)
    r += 1
    if context.get("branch_label"):
        ws.cell(row=r, column=1,
                value=f"{_('Branch')}: {context['branch_label']}").font = Font(size=9, italic=True)
        r += 1
    if context.get("period_label"):
        ws.cell(row=r, column=1, value=f"{_('Period')}: {context['period_label']}").font = Font(size=9, italic=True)
        r += 1
    r += 1

    if summary:
        ws.cell(row=r, column=1, value=str(_("Summary"))).font = Font(bold=True)
        r += 1
        for label, value in summary:
            ws.cell(row=r, column=1, value=str(label)).font = Font(bold=True)
            ws.cell(row=r, column=2, value=str(value))
            r += 1
        r += 1

    header_row = r
    for c, header in enumerate(_headers(report), start=1):
        cell = ws.cell(row=header_row, column=c, value=header)
        cell.font = Font(bold=True, color=black)
        cell.fill = PatternFill("solid", fgColor=brand_yellow)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    r += 1

    for row in rows:
        for c, (cell_value, col) in enumerate(zip(row, report.columns), start=1):
            # Write numbers as numbers so Excel can sum/sort them, rather
            # than as text — this is the main reason to offer xlsx at all.
            value = cell_value
            if col[1] == "r":
                try:
                    value = float(cell_value)
                except (TypeError, ValueError):
                    value = str(cell_value)
            else:
                value = str(cell_value)
            cell = ws.cell(row=r, column=c, value=value)
            cell.alignment = Alignment(horizontal="right" if col[1] == "r" else "left")
        r += 1

    # Freeze the header and auto-size columns to their content.
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    for c, header in enumerate(_headers(report), start=1):
        longest = max(
            [len(header)] + [len(str(row[c - 1])) for row in rows] or [len(header)]
        )
        ws.column_dimensions[get_column_letter(c)].width = min(max(longest + 3, 10), 45)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{_filename(report, "xlsx")}"'
    return response


# ---------------------------------------------------------------------------
# PDF (ReportLab)
# ---------------------------------------------------------------------------
def export_pdf(report, rows, summary, context):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
    )

    buffer = io.BytesIO()
    # Landscape: these reports are wide (up to 10 columns) and portrait
    # squeezes them unreadably.
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(A4),
        leftMargin=12 * mm, rightMargin=12 * mm,
        topMargin=12 * mm, bottomMargin=14 * mm,
        title=str(report.title), author=context["pharmacy_name"],
    )

    # Amharic needs a bundled Unicode TTF; Latin/digits stay in Helvetica.
    # See _register_pdf_fonts() and _rich() for why both are required.
    eth = _register_pdf_fonts()

    styles = getSampleStyleSheet()
    brand = colors.HexColor("#F1AB15")
    ink = colors.HexColor("#111111")
    muted = colors.HexColor("#666666")

    title_style = ParagraphStyle("t", parent=styles["Heading1"], fontSize=15,
                                 textColor=ink, spaceAfter=2, fontName="Helvetica-Bold")
    sub_style = ParagraphStyle("s", parent=styles["Normal"], fontSize=9,
                               textColor=muted)
    cell_style = ParagraphStyle("c", parent=styles["Normal"], fontSize=7.5,
                                leading=9.5)
    cell_right = ParagraphStyle("cr", parent=cell_style, alignment=TA_RIGHT)
    head_style = ParagraphStyle("h", parent=styles["Normal"], fontSize=7.5,
                                leading=9.5, textColor=ink, fontName="Helvetica-Bold")

    story = [
        Paragraph(_rich(context["pharmacy_name"], eth, bold=True), title_style),
        Paragraph(_rich(report.title, eth, bold=True), ParagraphStyle("t2", parent=styles["Heading2"],
                                                    fontSize=11, textColor=ink, spaceAfter=4,
                                                    fontName="Helvetica-Bold")),
        Paragraph(_rich(f"{_('Generated')}: {context['generated_at']}", eth), sub_style),
    ]
    if context.get("branch_label"):
        story.append(Paragraph(_rich(f"{_('Branch')}: {context['branch_label']}", eth), sub_style))
    if context.get("period_label"):
        story.append(Paragraph(_rich(f"{_('Period')}: {context['period_label']}", eth), sub_style))
    story.append(Spacer(1, 6 * mm))

    if summary:
        summary_data = [[Paragraph(_rich(l, eth, bold=True), head_style),
                         Paragraph(_rich(v, eth), cell_style)]
                        for l, v in summary]
        summary_table = Table(summary_data, hAlign="LEFT", colWidths=[55 * mm, 45 * mm])
        summary_table.setStyle(TableStyle([
            ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#DDDDDD")),
            ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#EEEEEE")),
            ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#FCEDCB")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        story += [summary_table, Spacer(1, 6 * mm)]

    if rows:
        table_data = [[Paragraph(_rich(h, eth, bold=True), head_style) for h in _headers(report)]]
        for row in rows:
            table_data.append([
                Paragraph(_rich(cell, eth), cell_right if col[1] == "r" else cell_style)
                for cell, col in zip(row, report.columns)
            ])
        table = Table(table_data, repeatRows=1, hAlign="LEFT")
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), brand),
            ("LINEBELOW", (0, 0), (-1, 0), 0.6, ink),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1),
             [colors.white, colors.HexColor("#FAFAFA")]),
            ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#E5E5E5")),
            ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#CCCCCC")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        story.append(table)
    else:
        story.append(Paragraph(_rich(_("No data for the selected filters."), eth), cell_style))

    def footer(canvas, doc_):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(muted)
        # Footer is drawn on the canvas (not a Paragraph), so it cannot use
        # inline font tags — keep it ASCII-safe.
        canvas.drawString(12 * mm, 8 * mm, context["footer_ascii"])
        canvas.drawRightString(
            landscape(A4)[0] - 12 * mm, 8 * mm, f"Page {doc_.page}"
        )
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    buffer.seek(0)
    response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="{_filename(report, "pdf")}"'
    return response


#: Unicode range for the Ethiopic script (Amharic). Used to decide which
#: runs of a string need the bundled TTF instead of built-in Helvetica.
_ETHIOPIC_RANGES = ((0x1200, 0x137F), (0x1380, 0x139F), (0x2D80, 0x2DDF))


def _is_ethiopic(ch):
    cp = ord(ch)
    return any(lo <= cp <= hi for lo, hi in _ETHIOPIC_RANGES)


def _register_pdf_fonts():
    """Registers Noto Sans Ethiopic for PDF output.

    Two things make this necessary and slightly fiddly:

    1. ReportLab's built-in Type 1 fonts (Helvetica et al.) have no Ethiopic
       glyphs, so Amharic silently renders as ■■■ boxes with no error.
    2. Noto Sans Ethiopic is an Ethiopic-only font — it contains no Latin
       letters or digits. Setting it as the document font would break every
       English label and every number instead.

    So neither font alone works: we register the Ethiopic TTF and then use
    `_rich()` to tag only the Ethiopic runs with it, leaving Latin and
    digits in Helvetica. Returns True if the TTF is available.
    """
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont as RLTTFont

    font_dir = Path(settings.BASE_DIR) / "static" / "fonts"
    candidates = [
        ("NotoEthiopic", font_dir / "NotoSansEthiopic-Regular.ttf"),
        ("NotoEthiopic-Bold", font_dir / "NotoSansEthiopic-Bold.ttf"),
    ]
    try:
        if not all(path.exists() for _name, path in candidates):
            logger.warning("Ethiopic PDF font missing; Amharic text may not render.")
            return False
        registered = pdfmetrics.getRegisteredFontNames()
        for name, path in candidates:
            if name not in registered:
                pdfmetrics.registerFont(RLTTFont(name, str(path)))
        return True
    except Exception:  # pragma: no cover - a font issue must never 500
        logger.warning("Could not register Ethiopic PDF font.", exc_info=True)
        return False


def _rich(value, ethiopic_available, bold=False):
    """Escapes `value` for ReportLab markup and wraps any Ethiopic runs in a
    <font> tag pointing at the bundled TTF, so mixed Amharic/English/numeric
    strings render correctly in a single Paragraph."""
    text = str(value)
    escaped = (
        text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    )
    if not ethiopic_available or not any(_is_ethiopic(c) for c in text):
        return escaped

    font = "NotoEthiopic-Bold" if bold else "NotoEthiopic"
    out, buf, in_eth = [], [], False
    for ch in escaped:
        eth = _is_ethiopic(ch)
        if eth != in_eth and buf:
            chunk = "".join(buf)
            out.append(f'<font name="{font}">{chunk}</font>' if in_eth else chunk)
            buf = []
        in_eth = eth
        buf.append(ch)
    if buf:
        chunk = "".join(buf)
        out.append(f'<font name="{font}">{chunk}</font>' if in_eth else chunk)
    return "".join(out)


EXPORTERS = {"csv": export_csv, "xlsx": export_xlsx, "pdf": export_pdf}